import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report_docx():
    src_file = "PAYENT_AGENT_PHASE_7_2_STAGING_DEPLOYMENT_REPORT.md"
    dst_file = "PAYENT_AGENT_PHASE_7_2_STAGING_DEPLOYMENT_REPORT.docx"
    
    if not os.path.exists(src_file):
        print(f"File {src_file} does not exist.")
        return

    doc = docx.Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(35, 35, 35)

    with open(src_file, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        
        parsed_rows = []
        for line in table_lines:
            parts = [c.strip() for c in line.strip().strip('|').split('|')]
            if all(set(c) <= set('-: ') for c in parts if c):
                continue
            parsed_rows.append(parts)

        if not parsed_rows:
            table_lines = []
            in_table = False
            return

        max_cols = max(len(r) for r in parsed_rows)
        tbl = doc.add_table(rows=len(parsed_rows), cols=max_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row in enumerate(parsed_rows):
            for c_idx, cell_value in enumerate(row):
                if c_idx < max_cols:
                    cell = tbl.cell(r_idx, c_idx)
                    cell.text = cell_value
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    
                    if r_idx == 0:
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
                        cell._tc.get_or_add_tcPr().append(shd)
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.font.size = Pt(9.5)
                    else:
                        for run in p.runs:
                            run.font.size = Pt(9.5)
                        if r_idx % 2 == 1:
                            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                            cell._tc.get_or_add_tcPr().append(shd)

        doc.add_paragraph()
        table_lines = []
        in_table = False

    for raw_line in lines:
        s = raw_line.rstrip()
        if s.startswith('|') and '|' in s[1:]:
            in_table = True
            table_lines.append(s)
            continue
        elif in_table:
            flush_table()

        if s.startswith('# '):
            p = doc.add_paragraph()
            run = p.add_run(s[2:])
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif s.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(s[3:])
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif s.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(s[4:])
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif s.startswith('- ') or s.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(s[2:])
            p.paragraph_format.space_after = Pt(2)
        elif s.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
        elif s.strip() == '':
            continue
        else:
            p = doc.add_paragraph()
            p.add_run(s)
            p.paragraph_format.space_after = Pt(4)

    if in_table:
        flush_table()

    doc.save(dst_file)
    print(f"Successfully generated {dst_file} (Size: {os.path.getsize(dst_file)} bytes)")

if __name__ == "__main__":
    create_report_docx()
